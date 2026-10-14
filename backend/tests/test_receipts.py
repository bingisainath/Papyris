"""Receipt scanning end to end with a fake AI reader (no network calls)."""

import uuid

import pytest
from sqlalchemy import select

from app.config.settings import settings
from app.models.ai import AppSetting, UserAISettings
from app.models.message import Message
from app.models.receipt import StoreDiscountRule
from app.services import secrets
from app.services.receipt_ai import pipeline
from app.services.receipt_ai.base import ExtractionError, ExtractionResult
from app.services.receipt_ai.schema import ReceiptExtraction
from tests.test_media import upload


def item(name, price, flags=(), category="other", quantity="1", confidence=0.97):
    return {"raw_text": f"{name.upper()} {price}", "name": name, "quantity": quantity, "unit": "each",
            "unit_price": None, "line_total": price, "category": category, "flags": list(flags),
            "tax_code": None, "confidence": confidence}


def adjustment(kind, label, amount, items=(), included=False):
    return {"kind": kind, "label": label, "amount": amount, "percent": None,
            "applies_to_items": list(items), "already_included_in_items": included}


TESCO = {
    "is_receipt": True, "store_name": "TESCO Ireland", "purchased_at": "2026-10-04", "currency": "EUR",
    "prices_include_tax": True,
    "items": [
        item("Milk", "1.80", category="dairy"),
        item("Cheese", "4.00", category="dairy"),
        item("Rice", "12.00", category="pantry"),
        item("Bread", "1.60", flags=["reduced"], category="bakery"),
        item("Croissants", "2.50", flags=["reduced"], category="bakery"),
    ],
    "adjustments": [
        adjustment("item_discount", "Reduced", "-1.20", [3]),
        adjustment("item_discount", "Reduced", "-2.50", [4]),
        adjustment("tax", "VAT incl.", "1.10", included=True),
    ],
    "subtotal": None, "total": "18.20", "payment_method": "card", "warnings": [],
}


class FakeExtractor:
    provider = "anthropic"

    def __init__(self):
        self.response = TESCO
        self.calls = []
        self.error = None

    async def extract(self, images, note, model, api_key):
        self.calls.append({"images": len(images), "model": model, "api_key": api_key})
        if self.error:
            raise ExtractionError(self.error)
        return ExtractionResult(ReceiptExtraction.model_validate(self.response), model, {"input_tokens": 1500, "output_tokens": 400})


@pytest.fixture
def fake_ai(monkeypatch):
    fake = FakeExtractor()
    monkeypatch.setitem(pipeline.EXTRACTORS, "anthropic", fake)
    monkeypatch.setattr(settings, "ANTHROPIC_API_KEY", "sk-ant-server-key")
    return fake


@pytest.fixture
def six(make_user, make_group):
    async def _make():
        people = [await make_user(f"p{i}") for i in range(1, 7)]
        group = await make_group(people[0], people[1:], "Flat")
        return people, group
    return _make


async def scan(client, user, conversation_id, note=None):
    url = (await upload(client, user)).json()["data"]["url"]
    r = await client.post(f"/api/v1/conversations/{conversation_id}/receipts",
                          json={"image_urls": [url], "note": note}, headers=user.headers)
    assert r.status_code == 200, r.text
    r = await client.get(f"/api/v1/receipts/{r.json()['data']['id']}", headers=user.headers)
    return r.json()["data"]


def review_body(receipt, assign: dict[int, list], enable_rule=False, payers=None):
    """Turn a receipt response into the PUT body the review screen sends."""
    return {
        "store_name": receipt["store_name"],
        "currency": receipt["currency"],
        "printed_total": "18.20",
        "items": [
            {"id": i["id"], "name": i["name"], "quantity": i["quantity"], "unit": i["unit"],
             "price": f"{i['gross_minor'] / 100:.2f}", "category": i["category"], "flags": i["flags"],
             "split_mode": "equal", "assignments": [{"user_id": u, "value": "1"} for u in assign[n]]}
            for n, i in enumerate(receipt["items"])
        ],
        "adjustments": [
            {"kind": a["kind"], "label": a["label"], "amount": f"{a['amount_minor'] / 100:.2f}", "percent": a["percent"],
             "item_indexes": a["item_indexes"], "allocation": a["allocation"], "assignee_ids": a["assignee_ids"],
             "source": a["source"], "enabled": a["enabled"] or (enable_rule and a["source"] == "store_rule")}
            for a in receipt["adjustments"]
        ],
        "payers": payers,
    }


async def test_scan_review_and_save_the_tesco_example(client, six, fake_ai, events, db):
    people, group = await six()
    p1, p2, p3, p4, p5, p6 = people
    r = await client.post("/api/v1/store-discounts", json={"store_name": "Tesco", "percent": "10"}, headers=p1.headers)
    assert r.status_code == 200, r.text

    receipt = await scan(client, p1, group, note="weekly shop")
    assert fake_ai.calls == [{"images": 1, "model": "claude-opus-5", "api_key": "sk-ant-server-key"}]
    assert receipt["status"] == "ready", receipt["warnings"]
    assert receipt["model"] == "claude-opus-5" and receipt["key_source"] == "app"
    assert [i["name"] for i in receipt["items"]] == ["Milk", "Cheese", "Rice", "Bread", "Croissants"]
    assert all(len(i["assignments"]) == 6 for i in receipt["items"])  # everyone by default
    assert receipt["totals"]["computed_total_minor"] == 1820 and receipt["totals"]["difference_minor"] == 0
    assert not any(a["kind"] == "tax" for a in receipt["adjustments"])  # VAT already inside prices

    rule = next(a for a in receipt["adjustments"] if a["source"] == "store_rule")
    assert rule["enabled"] is False  # not printed -> off by default
    assert rule["item_indexes"] == [0, 1, 2]  # not on reduced items
    assert rule["amount_minor"] == -178

    assign = {0: [p1.id], 1: [p2.id, p3.id], 2: [p.id for p in people], 3: [p4.id], 4: [p5.id]}
    r = await client.put(f"/api/v1/receipts/{receipt['id']}", json=review_body(receipt, assign, enable_rule=True), headers=p2.headers)
    assert r.status_code == 200, r.text
    reviewed = r.json()["data"]
    totals = {p["user_id"]: p["amount_minor"] for p in reviewed["totals"]["people"]}
    assert totals == {p1.id: 342, p2.id: 360, p3.id: 360, p4.id: 220, p5.id: 180, p6.id: 180}
    assert reviewed["totals"]["computed_total_minor"] == 1642
    assert reviewed["totals"]["difference_minor"] == 0  # own discount isn't on the receipt, so no mismatch
    assert reviewed["payers"] == [{"user_id": p1.id, "amount_minor": 1642}]

    events.clear()
    r = await client.post(f"/api/v1/receipts/{receipt['id']}/save", json={}, headers=p1.headers)
    assert r.status_code == 200, r.text
    expense = r.json()["data"]
    assert expense["source"] == "receipt" and expense["split_mode"] == "itemized"
    assert expense["total_minor"] == 1642 and expense["description"] == "TESCO Ireland"
    assert {s["user_id"]: s["amount_minor"] for s in expense["shares"]} == totals
    assert len(expense["receipt_images"]) == 1 and "sig=" in expense["receipt_images"][0]
    assert any(p.get("expenseId") == expense["id"] for _, p in events)

    card = (await db.execute(select(Message).where(Message.expense_id == uuid.UUID(expense["id"])))).scalar_one()
    assert "TESCO Ireland" in card.text

    balances = (await client.get(f"/api/v1/conversations/{group}/balances", headers=p2.headers)).json()["data"]
    eur = balances["currencies"][0]
    assert eur["currency"] == "EUR" and eur["my_net_minor"] == -360
    assert sum(d["amount_minor"] for d in eur["debts"]) == 1642 - 342
    assert all(d["to_user"] == p1.id for d in eur["debts"])

    # The next receipt remembers who milk is for
    again = await scan(client, p1, group)
    milk = again["items"][0]
    assert [a["user_id"] for a in milk["assignments"]] == [p1.id]
    assert milk["suggestion"].startswith("Usually for")


async def test_two_payers_and_mismatch_needs_review(client, six, fake_ai):
    people, group = await six()
    fake_ai.response = {**TESCO, "total": "19.00", "adjustments": TESCO["adjustments"][:2]}
    receipt = await scan(client, people[0], group)
    assert receipt["status"] == "needs_review"
    assert any("add up to" in w for w in receipt["warnings"])

    everyone = {n: [p.id for p in people] for n in range(5)}
    body = review_body(receipt, everyone, payers=[{"user_id": people[0].id, "amount": "10.00"}, {"user_id": people[1].id, "amount": "5.00"}])
    body["printed_total"] = "18.20"
    r = await client.put(f"/api/v1/receipts/{receipt['id']}", json=body, headers=people[0].headers)
    assert r.status_code == 200
    r = await client.post(f"/api/v1/receipts/{receipt['id']}/save", json={}, headers=people[0].headers)
    assert r.status_code == 400 and "€18.20" in r.json()["message"]

    body["payers"][1]["amount"] = "8.20"
    await client.put(f"/api/v1/receipts/{receipt['id']}", json=body, headers=people[0].headers)
    r = await client.post(f"/api/v1/receipts/{receipt['id']}/save", json={"description": "Big shop"}, headers=people[0].headers)
    assert r.status_code == 200, r.text
    payers = {p["user_id"]: p["amount_minor"] for p in r.json()["data"]["payers"]}
    assert payers == {people[0].id: 1000, people[1].id: 820}


async def test_single_other_payer_covers_the_whole_bill(client, six, fake_ai):
    people, group = await six()
    receipt = await scan(client, people[0], group)
    everyone = {n: [p.id for p in people] for n in range(5)}
    body = review_body(receipt, everyone, payers=[{"user_id": people[2].id}])
    r = await client.put(f"/api/v1/receipts/{receipt['id']}", json=body, headers=people[0].headers)
    assert r.json()["data"]["payers"] == [{"user_id": people[2].id, "amount_minor": 1820}]


async def test_unassigned_items_block_saving(client, six, fake_ai):
    people, group = await six()
    receipt = await scan(client, people[0], group)
    assign = {n: [people[0].id] for n in range(5)}
    assign[2] = []
    r = await client.put(f"/api/v1/receipts/{receipt['id']}", json=review_body(receipt, assign), headers=people[0].headers)
    assert r.json()["data"]["totals"]["unassigned_item_indexes"] == [2]
    r = await client.post(f"/api/v1/receipts/{receipt['id']}/save", json={}, headers=people[0].headers)
    assert r.status_code == 400 and "who each item is for" in r.json()["message"]


async def test_not_a_receipt_and_ai_errors(client, six, fake_ai):
    people, group = await six()
    fake_ai.response = {**TESCO, "is_receipt": False, "items": [], "adjustments": []}
    receipt = await scan(client, people[0], group)
    assert receipt["status"] == "failed" and "doesn't look like a receipt" in receipt["error"]

    fake_ai.error = "The Claude API key was rejected"
    receipt = await scan(client, people[0], group)
    assert receipt["status"] == "failed" and receipt["error"] == "The Claude API key was rejected"

    fake_ai.error = None
    fake_ai.response = TESCO
    r = await client.post(f"/api/v1/receipts/{receipt['id']}/retry", headers=people[1].headers)
    assert r.status_code == 403
    r = await client.post(f"/api/v1/receipts/{receipt['id']}/retry", headers=people[0].headers)
    assert r.status_code == 200
    r = await client.get(f"/api/v1/receipts/{receipt['id']}", headers=people[0].headers)
    assert r.json()["data"]["status"] == "ready"


@pytest.fixture
async def scan_limit_one(db):
    row = await db.get(AppSetting, pipeline.SCAN_LIMIT_KEY)
    previous, row.value = row.value, "1"
    await db.commit()
    yield
    row.value = previous
    await db.commit()


async def test_monthly_limit(client, make_user, make_dm, fake_ai, db, scan_limit_one):
    a, b = await make_user(), await make_user()
    dm = await make_dm(a, b)
    await scan(client, a, dm)
    url = (await upload(client, a)).json()["data"]["url"]
    r = await client.post(f"/api/v1/conversations/{dm}/receipts", json={"image_urls": [url]}, headers=a.headers)
    assert r.status_code == 429 and "reset on the 1st" in r.json()["message"]

    # A key saved before own keys were removed is ignored: the app's key and limit apply
    db.add(UserAISettings(user_id=uuid.UUID(a.id), anthropic_key_encrypted=secrets.encrypt("sk-ant-users-own-key")))
    await db.commit()
    r = await client.post(f"/api/v1/conversations/{dm}/receipts", json={"image_urls": [url]}, headers=a.headers)
    assert r.status_code == 429


async def test_admin_sets_scans_per_person(client, make_user, make_dm, fake_ai, db, scan_limit_one):
    """App limit 1: a tester gets a custom limit, someone else extra scans this month, a third is paused."""
    from sqlalchemy import update
    from app.models.user import User
    admin, tester, other = await make_user("admin"), await make_user(), await make_user()
    await db.execute(update(User).where(User.id == uuid.UUID(admin.id)).values(is_app_admin=True))
    await db.commit()
    dm = await make_dm(tester, other)
    base = f"/api/v1/admin/users"

    # Only admins, and a reason is required
    assert (await client.post(f"{base}/{tester.id}/scan-limit", json={"limit": 3, "reason": "beta"}, headers=other.headers)).status_code == 403
    assert (await client.post(f"{base}/{tester.id}/scan-limit", json={"limit": 3}, headers=admin.headers)).status_code == 422

    r = await client.post(f"{base}/{tester.id}/scan-limit", json={"limit": 3, "reason": "Beta tester"}, headers=admin.headers)
    assert r.status_code == 200 and r.json()["data"]["limit"] == 3 and r.json()["data"]["app_limit"] == 1
    for _ in range(3):
        await scan(client, tester, dm)
    url = (await upload(client, tester)).json()["data"]["url"]
    assert (await client.post(f"/api/v1/conversations/{dm}/receipts", json={"image_urls": [url]}, headers=tester.headers)).status_code == 429
    usage = (await client.get("/api/v1/ai/settings", headers=tester.headers)).json()["data"]["usage"]
    assert usage == {"scans_this_month": 3, "monthly_limit": 3}

    # Extra scans add up within the month, on top of the app's limit
    await scan(client, other, dm)
    for count in (1, 2):
        r = await client.post(f"{base}/{other.id}/extra-scans", json={"count": count, "reason": "A scan went wrong"}, headers=admin.headers)
    assert r.json()["data"] == {"limit": 4, "app_limit": 1, "custom_limit": None, "extra_this_month": 3, "used": 1}
    await scan(client, other, dm)

    # Extra scans from an earlier month don't count
    await db.execute(update(User).where(User.id == uuid.UUID(other.id)).values(extra_scans_month=pipeline._month_start().date().replace(year=2020)))
    await db.commit()
    details = (await client.get(f"{base}/{other.id}", headers=admin.headers)).json()["data"]["scans"]
    assert details["extra_this_month"] == 0 and details["limit"] == 1 and details["used"] == 2

    # 0 pauses scanning with its own message; null goes back to the app's limit
    await client.post(f"{base}/{tester.id}/scan-limit", json={"limit": 0, "reason": "Abuse"}, headers=admin.headers)
    r = await client.post(f"/api/v1/conversations/{dm}/receipts", json={"image_urls": [url]}, headers=tester.headers)
    assert r.status_code == 429 and "paused" in r.json()["message"]
    r = await client.post(f"{base}/{tester.id}/scan-limit", json={"limit": None, "reason": "Resolved"}, headers=admin.headers)
    assert r.json()["data"]["custom_limit"] is None and r.json()["data"]["limit"] == 1

    log = [a["action"] for a in (await client.get("/api/v1/admin/actions", headers=admin.headers)).json()["data"]]
    assert log.count("scan_limit_set") == 3 and log.count("extra_scans") == 2


async def test_no_server_key_and_access(client, make_user, make_dm, fake_ai, monkeypatch):
    monkeypatch.setattr(settings, "ANTHROPIC_API_KEY", "")
    a, b, outsider = await make_user(), await make_user(), await make_user()
    dm = await make_dm(a, b)
    url = (await upload(client, a)).json()["data"]["url"]
    r = await client.post(f"/api/v1/conversations/{dm}/receipts", json={"image_urls": [url]}, headers=a.headers)
    assert r.status_code == 503 and "isn't set up" in r.json()["message"]

    monkeypatch.setattr(settings, "ANTHROPIC_API_KEY", "sk-ant-server-key")
    receipt = await scan(client, a, dm)
    r = await client.get(f"/api/v1/receipts/{receipt['id']}", headers=outsider.headers)
    assert r.status_code == 404
    r = await client.post(f"/api/v1/conversations/{dm}/receipts",
                          json={"image_urls": ["https://evil.example/x.png"]}, headers=a.headers)
    assert r.status_code == 400


async def test_app_model_only_and_store_rule_matching(client, six, fake_ai, db):
    people, group = await six()
    owner = people[0]
    # Chats can't pick a model any more: the old setting is ignored and the app default is used
    from app.models.expense import ConversationSettings
    settings_row = await db.get(ConversationSettings, uuid.UUID(group))
    if settings_row is None:
        settings_row = ConversationSettings(conversation_id=uuid.UUID(group))
        db.add(settings_row)
    settings_row.receipt_model_id = 3
    db.add(StoreDiscountRule(owner_id=uuid.UUID(owner.id), store_name="Lidl", store_key="lidl", percent=5))
    await db.commit()
    receipt = await scan(client, owner, group)
    assert fake_ai.calls[-1]["model"] == "claude-opus-5"
    assert not any(a["source"] == "store_rule" for a in receipt["adjustments"])  # Lidl rule, Tesco receipt


def test_workspace_header_only_for_the_server_key(monkeypatch):
    import httpx2
    import anthropic
    from app.services.receipt_ai import claude

    monkeypatch.setattr(settings, "ANTHROPIC_API_KEY", "sk-ant-server")
    monkeypatch.setattr(settings, "ANTHROPIC_WORKSPACE_ID", "wrkspc_123")
    assert claude.make_client("sk-ant-server").default_headers.get("anthropic-workspace-id") == "wrkspc_123"
    assert "anthropic-workspace-id" not in claude.make_client("sk-ant-someone-else").default_headers

    request = httpx2.Request("POST", "https://api.anthropic.com/v1/messages")
    response = httpx2.Response(400, request=request)
    error = anthropic.BadRequestError("This API key is not scoped to a workspace, so this request must include the anthropic-workspace-id header", response=response, body=None)
    assert claude.is_workspace_error(error)


async def test_excel_export_summary_and_discounted_products(client, six, fake_ai, make_user):
    import io
    from openpyxl import load_workbook

    people, group = await six()
    p1, p2, p3, p4, p5, p6 = people
    receipt = await scan(client, p1, group)
    assign = {0: [p1.id], 1: [p2.id, p3.id], 2: [p.id for p in people], 3: [p4.id], 4: [p5.id]}
    r = await client.put(f"/api/v1/receipts/{receipt['id']}", json=review_body(receipt, assign), headers=p1.headers)
    assert r.status_code == 200, r.text
    assert (await client.post(f"/api/v1/receipts/{receipt['id']}/save", json={}, headers=p1.headers)).status_code == 200

    r = await client.get(f"/api/v1/conversations/{group}/expenses/export", headers=p2.headers)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="Flat-expenses-' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Split", "Expenses", "Discounts"]

    split = list(wb["Split"].iter_rows(values_only=True))
    people_cols = split[0][6:]
    assert len(people_cols) == 6
    head = split[1]
    assert head[1].startswith("TESCO Ireland · paid by") and head[5] == 18.2 and round(sum(v for v in head[6:] if v), 2) == 18.2
    rows = {r[1].strip(): r for r in split[2:]}
    assert list(rows) == ["Milk", "Cheese", "Rice", "Bread", "Croissants"]
    assert rows["Rice"][3:6] == (12, None, 12) and sorted(v for v in rows["Rice"][6:] if v) == [2] * 6
    assert sorted(v for v in rows["Cheese"][6:] if v) == [2, 2]
    assert rows["Bread"][3:6] == (1.6, -1.2, 0.4) and [v for v in rows["Bread"][6:] if v] == [0.4]
    for p_col in range(6, 12):  # each person's column adds up to their share in the expense row
        assert round(sum(r[p_col] or 0 for r in split[2:]), 2) == (head[p_col] or 0)

    summary = [[c for c in row if c is not None] for row in wb["Summary"].iter_rows(values_only=True)]
    assert summary[0] == ["Flat · expenses"]
    assert ["EUR · total spent", 18.2] in summary
    assert ["Saved with discounts", 3.7, "on 2 products (see Discounts)"] in summary
    assert "To settle up" in [row[0] for row in summary if row]

    expenses = list(wb["Expenses"].iter_rows(min_row=2, values_only=True))
    assert len(expenses) == 1 and expenses[0][1] == "TESCO Ireland" and expenses[0][3] == 18.2

    discounts = list(wb["Discounts"].iter_rows(min_row=2, values_only=True))
    assert [(d[2], d[3], d[4], d[5], d[7]) for d in discounts] == [
        ("Bread", 1.6, 0.4, 1.2, "Reduced"),
        ("Croissants", 2.5, 0, 2.5, "Reduced"),
    ]
    assert discounts[1][6] == 1.0  # 100% off

    outsider = await make_user("outsider")
    r = await client.get(f"/api/v1/conversations/{group}/expenses/export", headers=outsider.headers)
    assert r.status_code == 404
