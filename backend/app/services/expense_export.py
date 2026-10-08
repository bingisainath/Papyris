# backend/app/services/expense_export.py

"""
A chat's expenses as an Excel workbook, small and to the point:

  Summary    who paid what, each person's share and balance (per currency), how to settle up,
             and how much discounts saved
  Split      every expense with one column per person: for scanned receipts every product (price,
             discount, final price) and what each person pays for it, plus extras split separately
  Expenses   one row per expense: date, what, total, who paid, how it was split
  Discounts  every product that got cheaper on a scanned receipt: price before and after,
             what was saved (and %), which discount did it, and who the product was for
"""

import io
import uuid
from datetime import datetime, timezone
from decimal import Decimal
from types import SimpleNamespace

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.expense import Expense, Settlement
from app.models.receipt import Receipt
from app.models.user import User
from app.services.balances import net_balances, pairwise, simplify
from app.services.money import CURRENCIES, exponent, to_major
from app.services.receipt_ai.pipeline import DISCOUNT_KINDS, compute

HEADER = PatternFill("solid", fgColor="ECE8F5")
TITLE = Font(bold=True, size=14, color="443468")
BOLD = Font(bold=True)
MUTED = Font(color="64748B", size=10)
GOOD = Font(color="15803D")
BAD = Font(color="B91C1C")


def _fmt(currency: str) -> str:
    symbol = CURRENCIES.get(currency.upper(), (None, currency))[1]
    places = exponent(currency)
    return f'"{symbol}"#,##0{"." + "0" * places if places else ""};-"{symbol}"#,##0{"." + "0" * places if places else ""}'


def _money(minor: int, currency: str) -> Decimal:
    return to_major(minor, currency)


def _table(ws, row: int, headers: list[str]) -> int:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.font = BOLD
        cell.fill = HEADER
    return row + 1


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


async def build_workbook(db: AsyncSession, conversation_id: uuid.UUID, title: str, simplified: bool = True) -> bytes:
    expenses = (await db.execute(
        select(Expense).where(Expense.conversation_id == conversation_id, Expense.deleted_at.is_(None)).order_by(Expense.spent_at)
    )).scalars().all()
    settlements = (await db.execute(
        select(Settlement).where(Settlement.conversation_id == conversation_id, Settlement.deleted_at.is_(None))
    )).scalars().all()
    receipt_ids = [e.receipt_id for e in expenses if e.receipt_id]
    receipts = {r.id: r for r in (await db.execute(select(Receipt).where(Receipt.id.in_(receipt_ids)))).scalars().all()} if receipt_ids else {}

    user_ids = {p.user_id for e in expenses for p in e.payers} | {s.user_id for e in expenses for s in e.shares}
    user_ids |= {s.from_user for s in settlements} | {s.to_user for s in settlements}
    users = {u.id: (u.name or u.username) for u in (await db.execute(select(User).where(User.id.in_(user_ids)))).scalars().all()} if user_ids else {}
    name = lambda uid: users.get(uid if isinstance(uid, uuid.UUID) else uuid.UUID(str(uid)), "Someone")  # noqa: E731

    wb = Workbook()

    # ---------------- Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"{title} · expenses"
    ws["A1"].font = TITLE
    when = f"{expenses[0].spent_at:%d %b %Y} – {expenses[-1].spent_at:%d %b %Y}" if expenses else "No expenses yet"
    ws["A2"] = f"{when} · {len(expenses)} expense{'s' if len(expenses) != 1 else ''} · made {datetime.now(timezone.utc):%d %b %Y}"
    ws["A2"].font = MUTED
    row = 4

    plain = [SimpleNamespace(currency=e.currency, payers=[(str(p.user_id), p.amount_minor) for p in e.payers],
                             shares=[(str(s.user_id), s.amount_minor) for s in e.shares]) for e in expenses]
    paid_back = [SimpleNamespace(currency=s.currency, from_user=str(s.from_user), to_user=str(s.to_user), amount=s.amount_minor) for s in settlements]
    nets = net_balances(plain, paid_back)
    debts = simplify(nets) if simplified else pairwise(plain, paid_back)

    discount_rows = []  # filled below, also summed here
    splits = {}  # expense id -> (receipt, split result) for scanned receipts
    for e in expenses:
        receipt = receipts.get(e.receipt_id) if e.receipt_id else None
        if receipt is None:
            continue
        result, _ = compute(receipt)
        if result is None:
            continue
        splits[e.id] = (receipt, result)
        adjustments = {a.id: a for a in receipt.adjustments}
        for item in receipt.items:
            net = result.item_net.get(item.id)
            if net is None or net >= item.gross_minor:
                continue
            labels = [adjustments[a_id].label for a_id, amount in result.item_adjustments.get(item.id, [])
                      if a_id in adjustments and adjustments[a_id].kind in DISCOUNT_KINDS and amount]
            who = ", ".join(sorted({name(a.get("user_id")) for a in (item.assignments or []) if a.get("user_id")})) or "—"
            discount_rows.append((e.spent_at, receipt.store_name or e.description, item.name, item.gross_minor, net, e.currency, labels, who))

    for currency in sorted({e.currency for e in expenses}):
        fmt = _fmt(currency)
        paid: dict[str, int] = {}
        share: dict[str, int] = {}
        for e in expenses:
            if e.currency != currency:
                continue
            for p in e.payers:
                paid[str(p.user_id)] = paid.get(str(p.user_id), 0) + p.amount_minor
            for s in e.shares:
                share[str(s.user_id)] = share.get(str(s.user_id), 0) + s.amount_minor
        total = sum(e.total_minor for e in expenses if e.currency == currency)
        ws.cell(row=row, column=1, value=f"{currency} · total spent").font = BOLD
        c = ws.cell(row=row, column=2, value=_money(total, currency))
        c.number_format, c.font = fmt, BOLD
        row = _table(ws, row + 1, ["Person", "Paid", "Their share", "Balance", ""])
        people = sorted(set(paid) | set(share) | set(nets.get(currency, {})), key=lambda u: name(u).lower())
        for uid in people:
            balance = nets.get(currency, {}).get(uid, 0)
            ws.cell(row=row, column=1, value=name(uid))
            for col, minor in ((2, paid.get(uid, 0)), (3, share.get(uid, 0)), (4, balance)):
                cell = ws.cell(row=row, column=col, value=_money(minor, currency))
                cell.number_format = fmt
            ws.cell(row=row, column=4).font = GOOD if balance > 0 else BAD if balance < 0 else MUTED
            ws.cell(row=row, column=5, value="gets back" if balance > 0 else "owes" if balance < 0 else "settled").font = MUTED
            row += 1
        settle = [d for d in debts if d.currency == currency]
        if settle:
            row += 1
            ws.cell(row=row, column=1, value="To settle up").font = BOLD
            row += 1
            for d in settle:
                ws.cell(row=row, column=1, value=f"{name(d.from_user)} pays {name(d.to_user)}")
                cell = ws.cell(row=row, column=2, value=_money(d.amount, currency))
                cell.number_format = fmt
                row += 1
        saved = sum(r[3] - r[4] for r in discount_rows if r[5] == currency)
        if saved:
            row += 1
            ws.cell(row=row, column=1, value="Saved with discounts").font = BOLD
            cell = ws.cell(row=row, column=2, value=_money(saved, currency))
            cell.number_format, cell.font = fmt, GOOD
            ws.cell(row=row, column=3, value=f"on {sum(1 for r in discount_rows if r[5] == currency)} products (see Discounts)").font = MUTED
            row += 1
        row += 2
    _widths(ws, [28, 14, 14, 14, 12])

    # ---------------- Split: every item and who pays what for it
    ws = wb.create_sheet("Split")
    people = sorted({str(s.user_id) for e in expenses for s in e.shares if s.amount_minor}, key=lambda u: name(u).lower())
    headers = ["Date", "Expense / item", "Qty", "Price", "Discount", "Final"] + [name(u) for u in people]
    row = _table(ws, 1, headers)
    first_person = 7
    expense_fill = PatternFill("solid", fgColor="F6F4FB")
    for e in expenses:
        fmt = _fmt(e.currency)

        def put(r: int, col: int, minor: int | None, font=None) -> None:
            if minor is None:
                return
            cell = ws.cell(row=r, column=col, value=_money(minor, e.currency))
            cell.number_format = fmt
            if font:
                cell.font = font

        # The expense itself: total and each person's share of it
        payers = ", ".join(name(p.user_id) for p in e.payers if p.amount_minor)
        ws.cell(row=row, column=1, value=e.spent_at.replace(tzinfo=None)).number_format = "dd mmm yyyy"
        ws.cell(row=row, column=2, value=f"{e.description} · paid by {payers}" if payers else e.description)
        put(row, 6, e.total_minor, BOLD)
        shares = {str(s.user_id): s.amount_minor for s in e.shares}
        for n, uid in enumerate(people):
            put(row, first_person + n, shares.get(uid) or None, BOLD)
        for col in range(1, len(headers) + 1):
            ws.cell(row=row, column=col).fill = expense_fill
            if col != 1:
                ws.cell(row=row, column=col).font = BOLD
        row += 1

        if e.id not in splits:
            if e.split_mode != "equal":
                ws.cell(row=row, column=2, value=f"  split {e.split_mode}").font = MUTED
                row += 1
            continue
        receipt, result = splits[e.id]
        per_item: dict[int, dict[str, int]] = {}
        per_extra: dict[int, dict[str, int]] = {}
        for uid, parts in result.person_breakdown.items():
            for part in parts:
                target = per_item if part["type"] == "item" else per_extra
                target.setdefault(part["id"], {})[uid] = target.get(part["id"], {}).get(uid, 0) + part["amount"]
        for item in receipt.items:
            if item.id not in result.item_net:
                continue  # voided
            net = result.item_net[item.id]
            ws.cell(row=row, column=2, value=f"  {item.name}")
            quantity = Decimal(item.quantity or 1).normalize()
            ws.cell(row=row, column=3, value=int(quantity) if quantity == quantity.to_integral() else float(quantity))
            put(row, 4, item.gross_minor)
            put(row, 5, (net - item.gross_minor) or None, GOOD if net < item.gross_minor else None)
            put(row, 6, net)
            for n, uid in enumerate(people):
                put(row, first_person + n, per_item.get(item.id, {}).get(uid))
            row += 1
        labels = {a.id: f"{a.label or a.kind.replace('_', ' ').capitalize()} (split separately)" for a in receipt.adjustments}
        labels[None] = "Rounding to the cent"
        for adj_id, amounts in per_extra.items():
            ws.cell(row=row, column=2, value=f"  {labels.get(adj_id, 'Extra')}").font = MUTED
            put(row, 6, sum(amounts.values()))
            for n, uid in enumerate(people):
                put(row, first_person + n, amounts.get(uid))
            row += 1
    if not expenses:
        ws.cell(row=2, column=1, value="No expenses yet").font = MUTED
    ws.freeze_panes = "C2"
    _widths(ws, [13, 34, 6, 10, 10, 10] + [12] * len(people))

    # ---------------- Expenses
    ws = wb.create_sheet("Expenses")
    row = _table(ws, 1, ["Date", "What", "Category", "Total", "Paid by", "Split"])
    for e in expenses:
        fmt = _fmt(e.currency)
        ws.cell(row=row, column=1, value=e.spent_at.replace(tzinfo=None)).number_format = "dd mmm yyyy"
        ws.cell(row=row, column=2, value=e.description)
        ws.cell(row=row, column=3, value=(e.category or "other").capitalize())
        cell = ws.cell(row=row, column=4, value=_money(e.total_minor, e.currency))
        cell.number_format = fmt
        ws.cell(row=row, column=5, value=", ".join(name(p.user_id) for p in e.payers))
        ws.cell(row=row, column=6, value=", ".join(f"{name(s.user_id)} {_money(s.amount_minor, e.currency)}" for s in e.shares if s.amount_minor))
        row += 1
    ws.freeze_panes = "A2"
    _widths(ws, [13, 30, 13, 12, 22, 44])

    # ---------------- Discounts (which products got cheaper)
    ws = wb.create_sheet("Discounts")
    row = _table(ws, 1, ["Date", "Store", "Product", "Was", "Now", "Saved", "% off", "Discount", "For"])
    for spent_at, store, product, gross, net, currency, labels, who in sorted(discount_rows, key=lambda r: (r[0], r[2])):
        fmt = _fmt(currency)
        ws.cell(row=row, column=1, value=spent_at.replace(tzinfo=None)).number_format = "dd mmm yyyy"
        ws.cell(row=row, column=2, value=store)
        ws.cell(row=row, column=3, value=product)
        for col, minor in ((4, gross), (5, net), (6, gross - net)):
            ws.cell(row=row, column=col, value=_money(minor, currency)).number_format = fmt
        ws.cell(row=row, column=6).font = GOOD
        pct = ws.cell(row=row, column=7, value=float(Decimal(gross - net) / Decimal(gross)) if gross else 0)
        pct.number_format = "0%"
        ws.cell(row=row, column=8, value=", ".join(dict.fromkeys(labels)) or "Discount")
        ws.cell(row=row, column=9, value=who)
        row += 1
    if not discount_rows:
        ws.cell(row=2, column=1, value="No discounted products on scanned receipts yet").font = MUTED
    ws.freeze_panes = "A2"
    _widths(ws, [13, 20, 28, 10, 10, 10, 8, 24, 22])
    for sheet in wb.worksheets:
        for r in sheet.iter_rows():
            for cell in r:
                cell.alignment = Alignment(vertical="top")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
